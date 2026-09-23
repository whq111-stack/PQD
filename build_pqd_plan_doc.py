from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "docs"
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "基于物理条件扩散增强与暂态敏感双路径网络的复合电能质量扰动识别方法_完整方案.docx"

BLACK = "000000"
DARK_BLUE = "1F4E79"
MID_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
LIGHT_GRAY = "D9D9D9"
PALE_GRAY = "F4F6F8"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, name="Microsoft YaHei", size=10.5, color=BLACK, bold=False):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color="666666")


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, name="Cambria Math", size=11, color=BLACK)
    return p


def add_body(doc, text, first_line=True, space_after=5):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.55)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True, color=BLACK)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Cm(widths[i])
        set_cell_shading(cell, DARK_BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            if widths:
                cell.width = Cm(widths[i])
            set_cell_border(cell)
            set_cell_margins(cell, top=90, start=110, bottom=90, end=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_architecture_table(doc):
    rows = [
        ("数据输入", "原始一维电压序列 x ∈ R^L；默认 L = 640，扩展实验使用 L = 1280。", "不做 GAF、MTF、RGB 等二维图像转换"),
        ("局部暂态路径", "多尺度深度可分离空洞卷积 + 一阶差分残差", "脉冲、缺口、振荡起始和结束边界"),
        ("全局时序路径", "ModernTCN 主干；Mamba/SSM 作为替代对照", "跨周期幅值、谐波和闪变变化"),
        ("门控融合", "由局部突变强度、噪声估计和多尺度统计量生成融合门控", "不同工况下动态分配局部与全局特征"),
        ("组件解码器", "无序标签查询或扰动槽位解码器", "输出可组合扰动成分，而非固定组合类别"),
        ("任务头", "组件分类、机制分类、起止位置、严重程度、未知风险", "形成识别、定位和可靠性联合输出"),
        ("物理一致性约束", "成分重构、反事实一致性、跨域一致性", "抑制模型记忆仿真域，提升未见组合泛化"),
    ]
    return add_table(doc, ["模块", "实现", "作用"], rows, widths=[2.5, 8.1, 5.2], font_size=9.0)


def add_page_break(doc):
    doc.add_page_break()


def configure_document(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.1)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    styles = doc.styles
    set_style_font(styles["Normal"], size=10.5)
    body = styles["Body Text"]
    set_style_font(body, size=10.5)
    body.paragraph_format.line_spacing = 1.18
    body.paragraph_format.space_after = Pt(5)
    for level, size in ((1, 15), (2, 12), (3, 10.5)):
        st = styles[f"Heading {level}"]
        set_style_font(st, size=size, color=BLACK, bold=True)
        st.paragraph_format.keep_with_next = True
    title = styles["Title"]
    set_style_font(title, size=21, color=BLACK, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    subtitle = styles["Subtitle"]
    set_style_font(subtitle, size=12, color="555555")
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    for style_name in ("List Bullet", "List Number"):
        set_style_font(styles[style_name], size=10.5)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("复合电能质量扰动识别方法方案")
    set_run_font(r, size=8.5, color="666666")
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("基于物理条件扩散增强与暂态敏感双路径网络的复合电能质量扰动识别方法")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("完整论文方案与实验实施文档")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("研究对象：原始一维 PQD 序列的复合扰动识别、成分解耦与跨工况泛化")
    set_run_font(r, size=10.5, color="555555")

    add_heading(doc, "方案结论", 1)
    add_body(doc, "本方案保留现有 LGAN 项目的原始一维波形输入，不将信号转换为 GAF、MTF 或 RGB 图像。模型直接处理长度为 640 的一维电压序列，并在扩展实验中增加长度为 1280 的序列设置。最终方法由三个主体部分组成：物理条件扩散增强、暂态敏感双路径主干和组件级复合扰动解耦识别。三部分分别解决数据分布不足、短时暂态特征易丢失以及复合扰动组合泛化能力不足的问题。监督式早停不作为本论文内容。")
    add_body(doc, "论文的核心判断不是“更换一个最新网络即可提升准确率”，而是让模型学习可组合的扰动成分及其物理影响。数学模型负责提供可控的物理条件，扩散模型负责补足非理想波形分布，双路径主干负责同时观察局部突变与跨周期变化，组件级解码器负责输出扰动成分、机制、时间区间、严重程度和未知风险。")

    add_heading(doc, "目录结构", 1)
    toc_rows = [
        ("一", "研究定位与创新边界"),
        ("二", "数据集与物理条件扩散增强"),
        ("三", "暂态敏感双路径主干网络"),
        ("四", "组件级复合扰动解码与多任务学习"),
        ("五", "完整训练方法与实现细节"),
        ("六", "模型选择与对比策略"),
        ("七", "实验体系与评价指标"),
        ("八", "消融实验与论文结论判据"),
        ("九", "实施顺序、风险与备选方案"),
        ("十", "拟定论文结构与创新点表述"),
    ]
    add_table(doc, ["部分", "内容"], toc_rows, widths=[2.0, 13.8], font_size=9.5)

    add_page_break(doc)
    add_heading(doc, "一 研究定位与创新边界", 1)
    add_heading(doc, "1.1 现有代码和数据的事实基础", 2)
    add_body(doc, "现有仓库中的多标签数据由 640 个波形特征和 7 个扰动标签组成，数据文件包括 train_ml_300_rand、train_ml_100_rand 和 test_ml_100_rand。训练数据存在最多四个同时扰动的组合，标签组合数量为 48。数据读取代码将最后 7 列作为标签，并额外添加正常状态标签。现有 LGAN 的输入是形状为 batch × 1 × 640 的一维序列，主干包含一维卷积、标签解耦模块、CBAM 和 BRNN。")
    add_body(doc, "现有实现中存在三个与新论文直接相关的限制。第一，类别数、标签数和序列长度部分硬编码，难以扩展到带起止位置、机制标签和严重程度标签的新任务。第二，标签特征采用固定 one-hot，标签之间的可组合关系没有显式表达。第三，BRNN 的序列维度实际对应标签维度，不能替代对原始波形时间轴的建模。因此新方案需要保留数据接口的可复现性，同时重构主干和解码器。")

    add_heading(doc, "1.2 与参考论文的差异化边界", 2)
    boundary_rows = [
        ("逐点分类与时间定位", "已有论文已完成", "本方案仅将定位作为组件解耦的辅助任务，不把定位单独作为主创新"),
        ("条件网络和知识注入", "已有论文已完成", "本方案不使用扰动数量条件网络作为核心，而采用无序组件解码"),
        ("自适应小波和半监督", "已有论文已完成", "本方案不将小波或半监督作为主体，扩散模型服务于条件增强"),
        ("多尺度卷积与 Transformer", "已有论文已完成", "ModernTCN 或 SSM 只作为主干候选，创新在双路径和物理耦合"),
        ("标准数学数据集", "已有研究普遍采用", "本方案增加未见组合、反事实成分和设备域划分，改变实验问题"),
        ("本方案主创新", "未见于所给参考论文的完整组合", "物理条件扩散增强 + 暂态敏感双路径 + 组件级组合泛化"),
    ]
    add_table(doc, ["方向", "已有研究状态", "本方案处理方式"], boundary_rows, widths=[4.0, 4.4, 7.4], font_size=8.8)

    add_heading(doc, "1.3 研究问题", 2)
    add_body(doc, "论文围绕以下三个问题展开。问题一：仅使用标准数学仿真信号训练时，模型是否会依赖固定噪声、固定采样条件和固定扰动参数。问题二：训练阶段没有出现某种复合组合时，模型是否能够根据已学习的单一扰动成分识别该组合。问题三：脉冲、缺口和振荡等短时扰动在低信噪比下容易被池化和全局平均削弱，如何让主干网络保留这些特征并与跨周期信息共同建模。")
    add_body(doc, "由此，本文的实验目标不应只写成“提高总体准确率”，而应同时考察同分布识别、低信噪比鲁棒性、未见组合泛化、跨设备域泛化、暂态召回率和计算复杂度。")

    add_heading(doc, "1.4 最终创新点", 2)
    add_body(doc, "创新点一：提出物理条件一维扩散增强方法。将扰动类型、机制、幅值、持续时间、发生位置、噪声类型和设备域作为条件，在 IEEE Std. 1159-2019 约束的数学波形基础上生成具有更丰富非理想因素的训练样本，并通过物理一致性约束保证生成波形不会脱离设定的扰动规律。")
    add_body(doc, "创新点二：提出暂态敏感双路径主干网络。局部路径使用多尺度深度可分离空洞卷积和差分残差强化短时突变，全局路径使用 ModernTCN 或轻量状态空间模块建模跨周期变化，再由信号状态门控自适应融合两类特征。该主干直接处理一维序列，不需要二维图像化预处理。")
    add_body(doc, "创新点三：提出组件级复合扰动解耦识别方法。把一个复合 PQD 表示为无序扰动成分集合，联合预测具体类别、物理机制、起止位置和严重程度，并使用成分重构与反事实一致性损失约束特征学习，使模型从“记忆组合类别”转向“识别可组合成分”。")

    add_page_break(doc)
    add_heading(doc, "二 数据集与物理条件扩散增强", 1)
    add_heading(doc, "2.1 数据集总体策略", 2)
    add_body(doc, "采用“原始数据集基线 + 扩展数据集创新验证”的两级策略。原始 .mat 数据完全保留，用于复现 LGAN、比较新主干并证明方法在原任务上的有效性。扩展数据由可复现的数学函数生成，再由条件扩散模型学习仿真波形中的非理想变化，用于测试未见参数、未见组合和跨设备域泛化。")
    add_table(doc, ["数据层", "数据来源", "用途", "必须保持的条件"], [
        ("D0 原始基线", "仓库现有 .mat 文件", "复现和同分布对比", "640 点、原标签格式"),
        ("D1 参数扩展", "标准数学模型随机采样", "未见幅值、持续时间、频率", "标签和物理范围可追溯"),
        ("D2 反事实成分", "对复合信号逐成分增加或删除", "成分解耦和重构损失", "成分标签严格一致"),
        ("D3 环境域扩展", "噪声、传感器、采样和工况扰动", "跨域泛化和未知风险", "同一 PQD 标签跨域保持不变"),
        ("D4 组合泛化", "按组合划分训练和测试", "未见复合组合测试", "测试组合的单一成分见过"),
    ], widths=[3.0, 4.5, 5.0, 3.3], font_size=8.8)

    add_heading(doc, "2.2 数学扰动生成模型", 2)
    add_body(doc, "数学生成器不再只是增加样本数量，而是提供扩散模型可读取的物理条件和可验证标签。对每个样本记录基波、乘性扰动、加性扰动、噪声和设备变换。复合扰动可以写成如下形式：")
    add_formula(doc, "x(t) = x0(t) · M(t; cM) + A(t; cA) + n(t; cn, d)")
    add_body(doc, "其中 x0(t) 为带随机初相位和轻微工频偏移的基波，M(t) 表示暂降、暂升、中断和闪变等乘性变化，A(t) 表示谐波、缺口、脉冲和振荡等加性变化，n(t) 表示噪声，cM、cA、cn 和 d 分别为物理参数、噪声参数和设备域条件。每个扰动成分都保存起始点、结束点、幅值、频率、阻尼和相位等参数。")
    add_table(doc, ["扰动机制", "具体成分", "建议参数", "扩展测试"], [
        ("幅值变化", "暂降、暂升、中断", "幅值系数、持续周期、起止位置", "边界幅值和短持续时间"),
        ("波形畸变", "谐波、缺口", "谐波次数、相位、缺口深度", "时变谐波和不同相位"),
        ("暂态变化", "脉冲、振荡", "幅值、频率、阻尼、持续时间", "短时低幅值和高频振荡"),
        ("波动变化", "闪变", "调制深度、调制频率", "频率偏移和非平稳调制"),
        ("环境因素", "噪声和设备变换", "SNR、增益、偏置、量化、丢点", "未见噪声和跨设备"),
    ], widths=[3.0, 3.6, 5.4, 3.8], font_size=8.8)

    add_heading(doc, "2.3 扩散模型的使用方式", 2)
    add_body(doc, "推荐采用条件一维扩散模型作为数据增强器，而不是让扩散模型直接替代分类主干。扩散模型接收一个真实或数学生成的干净波形，通过前向过程加入噪声，再学习逆过程恢复波形。条件向量由标签、机制、物理参数、噪声类型和设备域编码组成。")
    add_formula(doc, "q(xk | x0) = N(sqrt(alphabar_k) x0, (1 - alphabar_k) I)")
    add_formula(doc, "eps_theta(xk, k, c) -> eps_hat")
    add_body(doc, "在实现上优先选择 DDPM 训练、DDIM 采样的一维 U-Net，因为它结构成熟、容易控制条件并且可以直接处理序列。若训练成本过高，可先使用轻量残差 1D U-Net，将扩散步嵌入、条件嵌入和局部卷积块组合。扩散模型生成的数据必须经过标签一致性筛选：用一个冻结的基础分类器和规则检测器检查生成波形的类别、频率、幅值和起止区间。")
    add_heading(doc, "2.4 物理条件和扩散损失", 2)
    add_body(doc, "普通扩散损失只要求预测噪声，无法保证生成结果满足 PQD 参数范围。因此增加物理一致性项。对生成样本计算频域能量、包络、局部突变、起止边界和成分检测结果，与条件参数对齐。")
    add_formula(doc, "Ldiff = Lnoise + lambda_phy Lphy + lambda_cond Lcond")
    add_body(doc, "Lnoise 为标准噪声预测损失；Lphy 约束幅值、频率、阻尼和起止位置；Lcond 约束生成波形经过冻结检测器后仍与条件标签一致。扩散模型不要求生成样本完全复制真实波形，而是要在物理有效范围内扩大环境和参数分布。")

    add_heading(doc, "2.5 数据划分和防止泄漏", 2)
    add_body(doc, "数据划分必须以扰动组合、参数范围和环境域为单位，而不能只对样本随机打乱。生成同一基础波形的多个噪声版本时，所有版本必须归入同一个划分，避免同源样本同时出现在训练集和测试集。")
    add_table(doc, ["测试协议", "训练可见内容", "测试内容", "论文要回答的问题"], [
        ("同分布", "全部类别和参数范围", "相同分布的新样本", "新模型是否保持基础性能"),
        ("参数外推", "参数范围 A", "未见参数范围 B", "是否学习到物理规律"),
        ("未见组合", "所有单一成分和部分组合", "未出现过的组合", "是否具备组合泛化能力"),
        ("跨噪声", "白噪声和部分 SNR", "彩色、脉冲和混合噪声", "是否依赖固定噪声"),
        ("跨设备", "设备域 1、2", "设备域 3、4", "是否依赖传感器特性"),
    ], widths=[3.0, 4.4, 4.4, 4.0], font_size=8.7)

    add_page_break(doc)
    add_heading(doc, "三 暂态敏感双路径主干网络", 1)
    add_heading(doc, "3.1 输入形式和总体结构", 2)
    add_body(doc, "TDSC 直接接收原始一维波形，输入张量为 batch × 1 × L。默认 L = 640，与现有数据一致；扩展实验可使用 L = 1280，但必须单独训练或采用长度无关的自适应池化。模型不生成二维图像，不使用 GAF、MTF、CWT 图像或 RGB 通道。")
    add_architecture_table(doc)
    add_body(doc, "TDSC 的设计依据是 PQD 同时具有短时突变和跨周期变化。局部路径保留高频和边界信息，全局路径利用大感受野和状态空间建模跨周期依赖，门控融合根据当前波形状态调整路径权重。")

    add_heading(doc, "3.2 局部暂态路径", 2)
    add_body(doc, "局部路径采用三个并行尺度的深度可分离空洞卷积块，建议卷积核大小为 3、5、7，膨胀率为 1、2、4。每个块包括深度卷积、逐点卷积、归一化、SiLU 激活和残差连接。为避免短脉冲在卷积前端被平滑，增加一阶差分输入：")
    add_formula(doc, "d(t) = x(t) - x(t - 1)")
    add_body(doc, "原始波形和差分波形分别经过浅层卷积后拼接。局部路径输出不使用过早的全局平均池化，而是保留时间位置特征，供后续组件定位头使用。")
    add_heading(doc, "3.3 全局时序路径和模型选择", 2)
    add_body(doc, "全局路径的首选骨干是 ModernTCN。它是面向通用时序分析的纯卷积结构，具有较大的有效感受野和相对清晰的工程实现，适合当前短一维序列。ModernTCN 不作为论文创新本身，而作为全局时序建模器嵌入 TDSC。")
    add_body(doc, "第二候选是选择性状态空间模型 Mamba。它可以作为替代骨干验证长程建模能力，但对 CUDA、版本和扫描算子依赖更敏感。若 Mamba 环境复现成本过高，不影响论文主线，因为 ModernTCN 已足以组成完整的双路径主干。标准 Transformer 只作为对比模型，不建议直接作为最终主干，原因是当前序列较短，标准自注意力的额外复杂度和数据需求不一定带来稳定收益。")
    add_table(doc, ["候选骨干", "使用位置", "优点", "风险", "建议"], [
        ("ModernTCN", "TDSC 全局路径", "官方实现、适合一维时序、实现稳定", "需要改造输入输出接口", "首选主干"),
        ("Mamba/SSM", "替代全局路径", "长程依赖、理论复杂度低", "环境和算子依赖", "第二候选或对照"),
        ("InceptionTime", "强 CNN 对照", "多尺度卷积、分类成熟", "不能单独体现跨路径门控", "建议加入基线"),
        ("TCN", "轻量对照", "简单、易部署", "全局表达能力有限", "加入基线"),
        ("标准 Transformer", "时序对照", "全局注意力直观", "数据量和计算量要求更高", "不作为主创新"),
        ("原始 CNN LGAN", "原始基线", "与现有工作一致", "暂态和组合建模不足", "必须保留"),
    ], widths=[3.1, 3.3, 4.0, 3.5, 2.1], font_size=8.3)

    add_heading(doc, "3.4 信号状态门控融合", 2)
    add_body(doc, "局部路径和全局路径不应简单相加。首先从输入波形提取轻量统计量，包括局部差分能量、峰度、短时方差、估计 SNR 和频带能量比；再由门控网络输出逐通道或逐时间段的融合权重。")
    add_formula(doc, "g = sigmoid(MLP([SNR_hat, E_delta, K_local, R_band]))")
    add_formula(doc, "F = g · F_local + (1 - g) · F_global")
    add_body(doc, "门控网络的作用不是提供额外的分类器，而是建立“信号状态—特征路径”的可解释联系。实验需要记录不同扰动、不同 SNR 下的平均门控值，检查脉冲和振荡是否倾向使用局部路径，闪变和谐波是否更多使用全局路径。")

    add_heading(doc, "3.5 主干参数建议", 2)
    add_table(doc, ["项目", "第一版设置", "调整范围", "选择原则"], [
        ("输入长度", "640", "640 或 1280", "先与原始数据一致"),
        ("局部通道", "32 -> 64 -> 128", "16 -> 256", "优先控制 FLOPs"),
        ("局部卷积尺度", "3、5、7", "3、5 或 3、7、11", "覆盖短脉冲和振荡周期"),
        ("膨胀率", "1、2、4", "1、2、4、8", "防止过大膨胀导致边界稀疏"),
        ("全局骨干", "ModernTCN 2 至 4 blocks", "1 至 6 blocks", "以验证集 Macro-F1 和延迟共同选择"),
        ("融合方式", "通道门控", "通道门控或时间门控", "先实现通道门控，后做扩展"),
        ("归一化", "BatchNorm 或 LayerNorm", "按 batch 大小选择", "小 batch 时优先 LayerNorm/GroupNorm"),
    ], widths=[3.0, 3.7, 4.5, 4.8], font_size=8.7)

    add_page_break(doc)
    add_heading(doc, "四 组件级复合扰动解码与多任务学习", 1)
    add_heading(doc, "4.1 复合扰动的集合表示", 2)
    add_body(doc, "复合扰动的本质是无序成分集合，例如 {暂降、谐波、振荡} 与 {振荡、暂降、谐波} 表示同一个物理事件。现有 LGAN 的标签引导模块将标签特征扩展后交给 BRNN，存在固定标签顺序和输出维度写死的问题。新方案使用类别查询向量或扰动槽位表示无序成分。")
    add_body(doc, "第一版推荐使用标签查询解码器。每个具体扰动类别拥有一个可学习查询 q_j，融合后的波形特征作为 Key 和 Value，通过交叉注意力得到类别证据。类别查询之间可以共享机制嵌入，机制嵌入由幅值、畸变、暂态和波动四类物理机制组成。这样既保持实现简单，又不依赖标签排列。")
    add_formula(doc, "z_j = CrossAttention(q_j + e_mech(j), F, F)")
    add_formula(doc, "p_j = sigmoid(MLP(z_j))")
    add_body(doc, "第二版可使用四个扰动槽位，每个槽位同时输出类别、存在概率、起止位置和严重程度，并用 Hungarian matching 完成预测槽位和真实成分的匹配。槽位方案更适合任意组合，但实现和调参成本高，建议在标签查询版本稳定后再尝试。")

    add_heading(doc, "4.2 输出任务", 2)
    add_table(doc, ["任务头", "输出", "监督标签", "作用"], [
        ("组件分类", "每类扰动概率 p_j", "7 类或扩展类别多标签", "识别具体扰动成分"),
        ("机制分类", "4 类机制概率", "幅值、畸变、暂态、波动", "增强物理层级表达"),
        ("时间定位", "起始和结束位置", "每个成分的 start/end", "描述扰动发生区间"),
        ("严重程度", "幅值、频率或阻尼回归", "数学模型参数", "输出可解释的扰动参数"),
        ("未知风险", "能量分数或证据不确定性", "未知组合和未知域测试", "避免对未知样本强制误分类"),
    ], widths=[3.0, 4.0, 4.3, 4.7], font_size=8.8)

    add_heading(doc, "4.3 成分重构分支", 2)
    add_body(doc, "为使组件特征具有物理含义，每个类别查询输出一个成分特征 z_j，并通过轻量重构器预测该扰动成分的波形贡献。对乘性扰动和加性扰动采用不同重构形式。")
    add_formula(doc, "x_hat = x0_hat · M_hat + sum_j A_hat_j")
    add_body(doc, "重构分支不要求精确恢复原始测量波形，而是要求去掉一个成分后，剩余波形与对应反事实样本一致。例如，输入为暂降加谐波，模型应能分别得到暂降成分和谐波成分；删除谐波成分的重构结果应接近只含暂降的数学波形。")

    add_heading(doc, "4.4 损失函数", 2)
    add_body(doc, "推荐总损失如下：")
    add_formula(doc, "L = L_cls + lambda_m L_mech + lambda_r L_rec + lambda_c L_cf + lambda_d L_domain + lambda_u L_unknown")
    add_table(doc, ["损失", "定义建议", "初始权重", "作用"], [
        ("L_cls", "加权 BCE 或 Asymmetric Loss", "1.0", "具体成分多标签识别"),
        ("L_mech", "四类机制 BCE/CE", "0.3", "物理层级监督"),
        ("L_rec", "波形 L1 + 频域损失", "0.2", "约束成分可重构"),
        ("L_cf", "同成分跨环境特征一致性", "0.1", "解耦环境变化"),
        ("L_domain", "域对抗或域分类反转", "0.05 至 0.1", "跨设备泛化"),
        ("L_unknown", "能量间隔或证据正则", "0.05", "开放集风险控制"),
    ], widths=[2.4, 6.1, 3.0, 5.0], font_size=8.8)
    add_body(doc, "若训练初期多任务损失互相干扰，先使用 L_cls + L_mech 训练分类模型，再逐步加入重构和一致性损失。权重不能只凭一次实验确定，应通过验证集 Macro-F1、未见组合 F1 和定位误差共同选择。")

    add_heading(doc, "4.5 与已有方法的实质差异", 2)
    add_body(doc, "本方案的组件解耦不是把四类机制各接一个分类头，也不是通过条件网络先预测扰动数量。关键差异在于：组件特征参与波形重构；训练时显式构造成分增删的反事实样本；测试时按未见组合划分数据；输出类别、区间和严重程度。只有同时具备这些设计，才可以将“物理一致性复合扰动解耦”作为创新点。")

    add_page_break(doc)
    add_heading(doc, "五 完整训练方法与实现细节", 1)
    add_heading(doc, "5.1 推荐训练阶段", 2)
    add_table(doc, ["阶段", "训练内容", "冻结模块", "输出"], [
        ("阶段 A", "在 D0 上训练原始 LGAN 和 TDSC 分类基线", "无", "基线性能和初始化模型"),
        ("阶段 B", "训练物理条件扩散模型", "数学生成器规则", "条件生成波形"),
        ("阶段 C", "用 D0 加 D1/D2 训练 TDSC 分类和机制头", "扩散模型", "可组合特征"),
        ("阶段 D", "加入重构、反事实一致性和定位损失联合微调", "可选冻结前端 5 至 10 个 epoch", "完整多任务模型"),
        ("阶段 E", "加入跨域和未知风险训练", "无", "跨工况可靠模型"),
    ], widths=[2.4, 6.2, 4.0, 3.9], font_size=8.7)
    add_body(doc, "建议采用分阶段训练，而不是一次性把所有模块同时打开。扩散生成器的质量需要先单独验证，分类模型需要先形成稳定的识别能力，之后再加入重构、一致性和未知风险任务。这样发生性能下降时可以定位具体模块。")

    add_heading(doc, "5.2 分类模型训练设置", 2)
    add_table(doc, ["设置项", "推荐初值", "备注"], [
        ("优化器", "AdamW", "比原始 Adam 更适合加入多任务正则"),
        ("学习率", "1e-3 主干，3e-4 解码头", "可使用 cosine decay 和 5 epoch warmup"),
        ("批大小", "64 至 288", "按显存调整，记录有效 batch"),
        ("训练 epoch", "100 至 200", "用早停保存最佳 Macro-F1；此处不是监督式早停创新"),
        ("分类损失", "加权 BCE 起步，Asymmetric Loss 对照", "按类别频率设置正负权重"),
        ("标签阈值", "验证集逐类校准", "不能固定认为 0.5 最优"),
        ("随机种子", "至少 3 个", "报告均值和标准差"),
        ("输入标准化", "按训练集统计量标准化", "禁止使用测试集统计量"),
    ], widths=[3.4, 4.0, 9.1], font_size=8.8)

    add_heading(doc, "5.3 代码重构要求", 2)
    add_body(doc, "新实现应移除硬编码的 num_classes、label_num、sequence_length 和 out.view(-1, 8)。建议建立一个配置文件，统一管理输入长度、类别名称、机制映射、最大成分数和数据路径。数据集读取器需要同时返回 waveform、component_labels、mechanism_labels、start_end、severity、snr、noise_type 和 domain_id。")
    add_body(doc, "推荐的模块文件划分如下：")
    add_table(doc, ["文件", "职责"], [
        ("data/pqd_generator.py", "标准数学模型、参数采样、组合划分和标签导出"),
        ("data/conditional_diffusion.py", "一维条件扩散训练和采样"),
        ("models/tdsc.py", "暂态敏感双路径主干和门控融合"),
        ("models/set_decoder.py", "标签查询或扰动槽位解码"),
        ("models/heads.py", "分类、机制、定位、严重程度和未知风险头"),
        ("losses/physical_losses.py", "重构、反事实、一致性和物理约束"),
        ("train_diffusion.py", "扩散模型训练"),
        ("train_tdsc.py", "识别模型分阶段训练"),
        ("evaluate.py", "统一指标、协议和结果保存"),
    ], widths=[5.0, 11.5], font_size=9.0)

    add_heading(doc, "5.4 训练和测试伪流程", 2)
    add_body(doc, "训练流程：读取数学条件和波形，生成或加载扩散增强样本，计算局部和全局路径特征，执行门控融合，使用组件查询解码，分别计算分类、机制、定位、严重程度和重构损失，最后更新模型。测试流程：严格按照测试协议构造输入，输出多标签结果和成分属性，计算同分布、低 SNR、未见组合、跨域和未知风险指标。")
    add_formula(doc, "F_local, F_global = TDSC(x, delta x)")
    add_formula(doc, "F = Gate(F_local, F_global, statistics(x))")
    add_formula(doc, "Y = Decoder(F, label_queries)")

    add_page_break(doc)
    add_heading(doc, "六 模型选择与对比策略", 1)
    add_heading(doc, "6.1 主模型选择结论", 2)
    add_body(doc, "最终主模型建议使用 ModernTCN 作为全局路径、深度可分离空洞卷积作为局部路径，形成 TDSC。选择 ModernTCN 的理由是：它有作者官方开源实现，面向通用时序分析，能够用纯一维结构扩大感受野，部署复杂度低于标准 Transformer。ModernTCN 只是主干组件，论文的创新必须写成“双路径、状态门控和组件级物理约束”，不能写成“首次将 ModernTCN 用于 PQD”。")
    add_body(doc, "Mamba/SSM 作为备选和对照。若在当前 Windows 和 PyTorch 环境中无法稳定安装或运行选择性扫描算子，直接使用 ModernTCN，不要为了追逐网络名称破坏实验可复现性。InceptionTime、TCN、CNN-LSTM、DCNN、ML-CNN 和原始 LGAN 构成基础对比集合。")

    add_heading(doc, "6.2 对比模型分组", 2)
    add_table(doc, ["分组", "模型", "比较目的"], [
        ("原始方法", "ML-CNN、DCNN、CNN-LSTM、原始 LGAN", "与仓库和原始论文保持可比"),
        ("参考方案", "师兄自适应小波专家网络的复现版本", "确认新方案相对已有改进的差异"),
        ("时序骨干", "TCN、InceptionTime、ModernTCN、Mamba/SSM", "比较现代一维时序建模器"),
        ("图像路线", "GAF/MTF + 轻量视觉模型，可选", "说明不使用二维转换的合理性"),
        ("新模型组件", "TDSC、TDSC + 组件解码、完整模型", "验证每项主体创新"),
    ], widths=[3.0, 7.0, 6.5], font_size=8.9)

    add_heading(doc, "6.3 不采用二维图像转换的理由", 2)
    add_body(doc, "本研究将原始波形直接作为一维序列进入 TDSC。第一，PQD 的起止时间、相位连续性和局部突变是原始时序中的直接信息，二维编码会引入额外表示过程。第二，当前目标包含定位、成分重构和严重程度估计，保留一维时间轴更利于建立这些监督。第三，二维路线已经在参考论文中用于 GAF、MTF 和 ViT，继续沿用会削弱方法差异。二维模型可以作为补充基线，但不进入主模型。")

    add_heading(doc, "6.4 开源实现使用规范", 2)
    add_body(doc, "ModernTCN 可参考作者官方实现 https://github.com/luodhhh/ModernTCN，Time-Series-Library 可用于统一比较不同现代时序模型 https://github.com/thuml/Time-Series-Library。使用开源代码时应记录 commit、许可证、依赖版本、输入适配修改和参数量，不能直接把仓库名称写成论文创新。")

    add_page_break(doc)
    add_heading(doc, "七 实验体系与评价指标", 1)
    add_heading(doc, "7.1 实验一 原始数据集复现", 2)
    add_body(doc, "使用现有 train_ml_300_rand、train_ml_100_rand 和 test_ml_100_rand，保持 640 点输入和原始标签定义。复现 ML-CNN、DCNN、CNN-LSTM、原始 LGAN，并训练 TDSC 和完整模型。该实验只回答基础识别性能，不作为组合泛化结论。")
    add_table(doc, ["指标类别", "指标", "报告方式"], [
        ("多标签识别", "Exact-match accuracy、Hamming loss", "与仓库结果保持可比"),
        ("类别均衡", "Macro-F1、Micro-F1、Macro-Recall", "必须报告，防止总体准确率掩盖暂态漏检"),
        ("排序性能", "Average Precision、Ranking Loss、Coverage", "保持与原 metrics.py 一致并补充解释"),
        ("复杂度", "Params、FLOPs、CPU/GPU 推理时间", "统一输入长度和硬件"),
    ], widths=[4.0, 6.2, 6.3], font_size=8.9)

    add_heading(doc, "7.2 实验二 噪声鲁棒性", 2)
    add_body(doc, "测试无噪声、50 dB、40 dB、30 dB、20 dB 和 10 dB。除高斯白噪声外，增加彩色噪声、脉冲噪声和混合噪声。重点报告暂态脉冲、振荡和缺口的 Recall、Macro-F1 及性能下降幅度。")
    add_formula(doc, "RobustnessDrop = F1_clean - F1_20dB")

    add_heading(doc, "7.3 实验三 未见复合组合泛化", 2)
    add_body(doc, "按组合划分训练集和测试集。训练集必须包含所有待测单一成分，但隐藏部分双重或三重组合。例如训练阶段出现暂降、谐波、振荡、暂降加谐波，测试阶段使用暂降加振荡、谐波加振荡和暂降加谐波加振荡。该实验是物理成分解耦创新的核心证据。")
    add_formula(doc, "CompositionGap = F1_seen - F1_unseen")
    add_body(doc, "如果完整模型只在随机划分测试集上提升，而在未见组合测试中没有缩小 CompositionGap，则不能声称实现了组合泛化。")

    add_heading(doc, "7.4 实验四 未见参数和边界条件", 2)
    add_table(doc, ["因素", "训练范围示例", "测试范围示例"], [
        ("暂降深度", "0.2 至 0.6", "0.1 至 0.2、0.6 至 0.9"),
        ("持续时间", "1 至 5 个周期", "0.5 至 1 个周期、5 至 9 个周期"),
        ("振荡频率", "100 至 300 Hz", "300 至 500 Hz"),
        ("发生位置", "窗口中部为主", "窗口开头、结尾和边界"),
        ("工频", "50 Hz 附近固定", "49.5 至 50.5 Hz"),
        ("传感器", "固定增益和量化", "增益、偏置、削顶、丢点"),
    ], widths=[4.0, 6.0, 6.5], font_size=8.8)

    add_heading(doc, "7.5 实验五 跨噪声和跨设备域", 2)
    add_body(doc, "训练域和测试域使用不同噪声类型、增益偏置、采样抖动、量化精度、基线漂移和工频偏移。报告域内 Macro-F1、跨域 Macro-F1、性能下降、暂态 Recall 和特征可视化。若有真实或半实物波形，应将其作为独立域测试，不能与仿真样本随机混合。")

    add_heading(doc, "7.6 实验六 定位和严重程度估计", 2)
    add_table(doc, ["任务", "指标", "重点场景"], [
        ("起止定位", "Start MAE、End MAE、区间 IoU", "短暂态、窗口边界、强噪声"),
        ("持续时间", "Duration MAE 和相对误差", "不同持续周期"),
        ("严重程度", "幅值/频率/阻尼 RMSE 或 MAE", "未见参数范围"),
        ("成分重构", "时域 L1、频域误差、相关系数", "成分增删反事实样本"),
    ], widths=[3.5, 6.1, 6.9], font_size=8.8)
    add_body(doc, "定位不是本论文单独的主创新，因为参考论文已经完成逐点分类和时间定位。本方案使用定位和严重程度作为组件解耦的辅助任务，证明成分特征具有可解释的时间和物理属性。")

    add_heading(doc, "7.7 实验七 未知扰动和未知组合拒识", 2)
    add_body(doc, "测试集中加入训练阶段没有出现的组合、噪声类型、振荡频率和人为构造的新型波形。模型除输出已知类别概率外，还要输出未知风险。评价 AUROC、AUPR、FPR95、Unknown Recall、ECE 和 Risk-Coverage 曲线。")
    add_body(doc, "如果实现复杂，第一版可采用能量分数或 MC Dropout；如果要加强理论完整性，再加入 evidential head 或 conformal prediction。未知检测不应改变已知类上的基本分类评价。")

    add_page_break(doc)
    add_heading(doc, "八 消融实验与论文结论判据", 1)
    add_heading(doc, "8.1 主体消融矩阵", 2)
    add_table(doc, ["模型", "扩散增强", "TDSC", "组件解码", "重构/反事实", "跨域/未知"], [
        ("B0 原始 LGAN", "×", "×", "×", "×", "×"),
        ("B1", "×", "√", "×", "×", "×"),
        ("B2", "√", "√", "×", "×", "×"),
        ("B3", "√", "√", "√", "×", "×"),
        ("B4", "√", "√", "√", "√", "×"),
        ("B5 完整模型", "√", "√", "√", "√", "√"),
    ], widths=[3.7, 2.5, 2.5, 2.8, 3.3, 2.7], font_size=9.0)
    add_body(doc, "每组至少在原始同分布、20 dB、未见组合和跨域四种协议下报告结果。只在同分布测试中做一组消融不足以说明各模块的作用。")

    add_heading(doc, "8.2 主干细化消融", 2)
    add_table(doc, ["实验", "改变项", "预期回答"], [
        ("C1", "仅局部路径", "局部特征对暂态的贡献"),
        ("C2", "仅全局路径", "跨周期特征的贡献"),
        ("C3", "简单相加替代门控", "门控融合是否有效"),
        ("C4", "去掉差分残差", "差分输入是否提升短时扰动召回"),
        ("C5", "ModernTCN 替换为 TCN/SSM", "骨干选择的影响"),
        ("C6", "不同局部卷积尺度", "暂态尺度覆盖的影响"),
    ], widths=[2.8, 7.3, 5.4], font_size=8.9)

    add_heading(doc, "8.3 扩散增强细化消融", 2)
    add_table(doc, ["实验", "增强方式", "评价重点"], [
        ("D1", "无增强", "基础性能"),
        ("D2", "数学随机参数增强", "传统增强收益"),
        ("D3", "普通 1D 扩散增强", "扩散模型本身收益"),
        ("D4", "条件扩散无物理损失", "条件控制作用"),
        ("D5", "物理条件扩散", "物理约束和组合泛化"),
        ("D6", "只使用扩散样本训练", "检查生成数据是否引入偏差"),
    ], widths=[2.8, 7.3, 5.4], font_size=8.9)

    add_heading(doc, "8.4 可视化和诊断实验", 2)
    add_bullet(doc, "绘制局部路径、全局路径和门控融合的 Grad-CAM 或输入梯度，检查模型关注区域是否落在扰动发生区间。")
    add_bullet(doc, "绘制同一扰动在不同噪声和设备域下的特征分布，比较原始 LGAN 与完整模型的域间聚集情况。")
    add_bullet(doc, "展示复合波形、真实成分、模型成分重构和删除某一成分后的反事实波形。")
    add_bullet(doc, "展示暂态类别在不同 SNR 下的召回率曲线，不以总体准确率替代暂态分析。")

    add_heading(doc, "8.5 结果判据", 2)
    add_body(doc, "建议在论文开始前设定结果判据，避免只挑选有利实验。完整模型至少应满足：原始数据集性能不低于原始 LGAN；20 dB 条件下暂态 Macro-Recall 有稳定提升；未见组合的 F1 明显高于普通标签分类器；跨域性能下降幅度低于基线；模型复杂度和推理时间可报告且没有失控。若某个目标未达到，应如实将该模块降为探索性结果，而不是保留为核心结论。")

    add_page_break(doc)
    add_heading(doc, "九 实施顺序 风险与备选方案", 1)
    add_heading(doc, "9.1 分阶段实施计划", 2)
    add_table(doc, ["阶段", "任务", "完成标志"], [
        ("第一阶段", "清理输入长度、标签数和路径硬编码；复现原始 LGAN", "基线指标和训练曲线可复现"),
        ("第二阶段", "实现 TDSC，完成局部路径、ModernTCN 全局路径和门控融合", "TDSC 在 D0 上稳定训练"),
        ("第三阶段", "实现数学数据生成器和未见组合划分", "可导出 waveform 和完整元数据"),
        ("第四阶段", "实现组件查询解码、定位和严重程度头", "完成多任务模型"),
        ("第五阶段", "实现 1D 条件扩散增强和物理筛选", "生成样本通过物理和标签检查"),
        ("第六阶段", "加入重构、反事实和跨域损失", "完成主体消融实验"),
        ("第七阶段", "完成未知拒识、真实/半实物测试和论文图表", "形成可投稿实验闭环"),
    ], widths=[3.1, 8.7, 3.7], font_size=8.8)

    add_heading(doc, "9.2 主要风险和处理", 2)
    add_table(doc, ["风险", "表现", "处理"], [
        ("扩散样本质量不足", "生成波形标签不一致或不符合物理范围", "先用数学规则和冻结分类器筛选；降低扩散模块权重"),
        ("多任务冲突", "分类提升但定位或重构不稳定", "分阶段加损失；使用梯度裁剪和权重搜索"),
        ("未见组合无提升", "随机测试提升，组合测试下降", "强化成分增删反事实和组合划分，检查数据泄漏"),
        ("ModernTCN 改造困难", "输入输出维度不匹配", "先用 TCN/InceptionTime 验证 TDSC 逻辑，再替换骨干"),
        ("Mamba 环境问题", "CUDA 或依赖报错", "只作为可选对照，不阻塞主模型"),
        ("真实数据不足", "无法获得大规模实测标签", "保留跨域仿真和少量半实物测试，明确局限性"),
    ], widths=[3.8, 5.4, 6.3], font_size=8.7)

    add_heading(doc, "9.3 最小可发表版本", 2)
    add_body(doc, "如果时间或算力有限，最小版本包括：TDSC 双路径主干、标签查询组件解码、数学模型生成的未见组合测试、起止位置辅助任务和低 SNR 实验。物理条件扩散作为增强模块加入后，论文完整性明显提高；跨域未知拒识可作为第二阶段扩展。不要为了加入 Mamba、标准 Transformer 或复杂开放集方法而牺牲主实验的可重复性。")

    add_heading(doc, "9.4 不建议加入的内容", 2)
    add_bullet(doc, "不加入监督式早停。它只能作为普通工程优化，技术贡献不足，且会分散组合泛化主线。")
    add_bullet(doc, "不把二维图像转换作为主路线。参考论文已经使用 GAF、MTF 和 ViT，且不利于直接定位和一维重构。")
    add_bullet(doc, "不把小波、Transformer、Mamba 或 ModernTCN 的名称单独写成创新。它们是组件或基线，创新来自问题定义和物理约束。")
    add_bullet(doc, "不在同一随机样本划分中同时声称实现跨域和未见组合泛化。两个测试协议必须独立设计。")

    add_page_break(doc)
    add_heading(doc, "十 拟定论文结构与创新点表述", 1)
    add_heading(doc, "10.1 论文结构", 2)
    add_table(doc, ["章节", "主要内容"], [
        ("第一章 绪论", "PQD 复杂化、复合组合、暂态识别和仿真到现场泛化问题；研究贡献"),
        ("第二章 PQD 数学模型与跨域数据集", "标准数学模型、成分标签、反事实样本、未见组合划分和设备域"),
        ("第三章 物理条件扩散增强", "一维条件扩散、物理约束、样本筛选和生成质量评价"),
        ("第四章 暂态敏感双路径网络", "局部暂态路径、ModernTCN 全局路径、状态门控和组件查询解码"),
        ("第五章 物理一致性多任务识别", "分类、机制、定位、严重程度、重构和一致性损失"),
        ("第六章 实验分析", "同分布、低 SNR、未见组合、跨域、未知风险和消融"),
        ("第七章 结论与展望", "贡献、局限性和现场数据拓展"),
    ], widths=[4.2, 11.3], font_size=8.9)

    add_heading(doc, "10.2 建议的创新点正式表述", 2)
    add_body(doc, "第一，提出一种物理条件一维扩散增强方法，将复合电能质量扰动的类型、机制、时序参数、环境噪声和设备域作为条件，生成物理参数可控的非理想扰动样本，并通过物理一致性损失和标签筛选提升增强数据的有效性。")
    add_body(doc, "第二，提出一种暂态敏感双路径时序主干网络。局部路径采用多尺度深度可分离空洞卷积和差分残差强化短时突变，全局路径采用 ModernTCN 或轻量状态空间模块建模跨周期变化，并通过信号状态门控实现局部与全局特征的自适应融合。")
    add_body(doc, "第三，提出一种面向复合扰动的组件级无序解码和物理一致性多任务学习方法，将复合扰动表示为可组合的扰动成分集合，联合学习成分类别、机制、起止位置和严重程度，并利用成分重构与反事实一致性约束提升未见复合组合下的识别能力。")
    add_body(doc, "第四，构建面向组合泛化和跨工况泛化的评价体系，通过未见组合、未见参数、跨噪声、跨设备和未知扰动测试，验证模型从仿真分布向复杂应用条件迁移的可靠性。")

    add_heading(doc, "10.3 题目和摘要中的表达边界", 2)
    add_body(doc, "论文题目采用用户确定的名称：基于物理条件扩散增强与暂态敏感双路径网络的复合电能质量扰动识别方法。摘要中应突出“原始一维序列”“物理条件扩散”“双路径局部—全局建模”“组件级组合泛化”四个关键词，不把二维图像、监督式早停或强化学习写入主体方法。")

    add_heading(doc, "参考资料与开源实现", 1)
    refs = [
        "[1] 现有 LGAN 项目代码及数据集，D:\\code\\LGAN-PQD-main。",
        "[2] 刘宇龙等. 基于多级注意力机制融合的电能质量扰动点分类及时间定位方法研究. 中国电机工程学报, 2024, 44(11): 4298-4309。",
        "[3] 黄杰等. 基于条件网络和知识注入的多标签学习电能质量扰动识别与分类方法. 2025。",
        "[4] 徐玉珍等. 融合时频自适应小波卷积网络的半监督电能质量扰动识别框架. 中国电机工程学报, 2026, 46(8): 3118-3129。",
        "[5] 李贝奥. 基于 Transformer 的复合电能质量扰动识别研究. 华中科技大学硕士学位论文, 2024。",
        "[6] 高鑫. 基于数据驱动的微电网线路故障诊断与电能质量扰动识别方法研究. 学位论文。",
        "[7] ModernTCN 官方实现: https://github.com/luodhhh/ModernTCN。",
        "[8] Time-Series-Library: https://github.com/thuml/Time-Series-Library。",
        "[9] IEEE Std. 1159-2019, IEEE Recommended Practice for Monitoring Electric Power Quality。",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(ref)
        set_run_font(r, size=9.5)

    add_heading(doc, "执行摘要", 1)
    add_body(doc, "最终实施建议可以概括为：原始一维波形直接输入 TDSC；数学函数负责生成有物理标签的基础样本和未见组合；一维条件扩散负责模拟非理想环境；双路径主干分别捕获暂态局部特征和跨周期全局特征；组件级解码器输出复合扰动成分；重构和反事实一致性损失确保成分具有物理意义；实验以未见组合、低 SNR 和跨工况泛化为主要证据。")
    add_body(doc, "这条路线同时改变了数据、主干和任务定义，能够避免论文仅被评价为“在 LGAN 上替换模块”。")

    doc.core_properties.title = "基于物理条件扩散增强与暂态敏感双路径网络的复合电能质量扰动识别方法"
    doc.core_properties.subject = "论文方案与实验实施"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
